"""Qarkullimet Ditore: daily cash book built from existing sales/purchases/expenses/payments/returns."""
import io
from datetime import datetime, date as Date, timedelta
from typing import Optional, Literal, Dict
from zoneinfo import ZoneInfo

from bson import ObjectId
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from finance import get_rates

TZ = ZoneInfo("Europe/Belgrade")
DENOMS = ["500", "200", "100", "50", "20", "10", "5", "2", "1", "0.5", "0.2", "0.1"]
OVERRIDABLE = ("hyrje", "dalje", "method", "description", "invoice_no", "cost")


def r2(x) -> float:
    return round(float(x or 0) + 0.0, 2)


def local_day(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso).astimezone(TZ).date().isoformat()
    except (TypeError, ValueError):
        return ""


def parse_day(s: str) -> Date:
    try:
        return Date.fromisoformat(s)
    except (TypeError, ValueError):
        raise HTTPException(400, "Data nuk është e vlefshme")


class ManualIn(BaseModel):
    date: str
    kind: Literal["in", "out"]
    amount: float
    method: Literal["cash", "bank"] = "cash"
    description: str = ""
    invoice_no: str = ""
    in_profit: bool = False


class OverrideIn(BaseModel):
    key: str
    hyrje: Optional[float] = None
    dalje: Optional[float] = None
    method: Optional[Literal["cash", "bank"]] = None
    description: Optional[str] = None
    invoice_no: Optional[str] = None
    cost: Optional[float] = None
    reason: str = ""


class SettingsIn(BaseModel):
    start_date: str
    opening_cash: float


class CountIn(BaseModel):
    counts: Dict[str, int]
    note: str = ""


class DeleteIn(BaseModel):
    reason: str


def make_router(db, require_admin, now_utc, company_header):
    router = APIRouter(prefix="/daybook")

    async def audit(key: str, action: str, before, after, user: dict, reason: str = ""):
        await db.daybook_audit.insert_one({
            "key": key, "action": action, "before": before, "after": after, "reason": reason,
            "user_id": user["id"], "user_name": user.get("name") or user.get("email", ""), "at": now_utc().isoformat(),
        })

    async def settings() -> dict:
        s = await db.daybook_settings.find_one({"_id": "cash"})
        return {"start_date": (s or {}).get("start_date"), "opening_cash": r2((s or {}).get("opening_cash", 0)), "set": bool(s)}

    async def collect(d_from: str, d_to: str) -> list:
        lo = (parse_day(d_from) - timedelta(days=1)).isoformat()
        hi = (parse_day(d_to) + timedelta(days=2)).isoformat()
        in_range = lambda d: d_from <= d <= d_to  # noqa: E731
        created = {"created_at": {"$gte": lo, "$lt": hi}}
        customers = {str(c["_id"]): c.get("name", "") async for c in db.customers.find({}, {"name": 1})}
        suppliers = {str(c["_id"]): c.get("name", "") async for c in db.suppliers.find({}, {"name": 1})}
        overrides = {o["key"]: o async for o in db.daybook_overrides.find({})}
        out = []

        sales_q = {"$or": [{"document_date": {"$gte": d_from, "$lte": d_to}},
                           {"document_date": {"$in": [None, ""]}, **created}]}
        sales_by_id = {}
        async for s in db.sales.find(sales_q):
            sales_by_id[str(s["_id"])] = s
            d = s.get("document_date") or local_day(s.get("created_at"))
            if not in_range(d):
                continue
            items = s.get("items", []) or []
            missing = any(float(i.get("cost_price") or 0) <= 0 for i in items)
            cost = sum(float(i.get("cost_price") or 0) * int(i.get("quantity") or 0) for i in items)
            names = ", ".join(i.get("product_name", "") for i in items)
            cust = s.get("customer_name") or "Qytetar"
            pm = s.get("payment_method", "cash")
            paid = r2(s.get("paid_amount"))
            if "card_paid" in s:
                cash_part, bank_part = r2(s.get("cash_paid")), r2(s.get("card_paid"))
                if pm == "debt":
                    cash_part = paid
            else:  # sales saved before the Card-Bank update
                cash_part, bank_part = (0.0, paid) if pm == "card" else (paid, 0.0)
            out.append({
                "key": f"sale:{s['_id']}", "source": "sale", "source_id": str(s["_id"]), "date": d,
                "time": s.get("created_at", ""), "type": "Shitje", "description": f"{names} · {cust}",
                "invoice_no": s.get("invoice_no", ""), "hyrje": r2(s.get("paid_amount")), "dalje": 0.0,
                "method": "mixed" if cash_part and bank_part else ("bank" if bank_part else "cash"),
                "cash_in": cash_part, "bank_in": bank_part, "vat": r2(s.get("vat")), "bank_fee": r2(s.get("bank_fee")),
                "revenue": r2(s.get("total")), "cost": r2(cost),
                "cost_missing": missing, "opex": 0.0, "value": r2(s.get("total")),
                "note": "Me borxh" if pm == "debt" and float(s.get("due_amount") or 0) > 0 else "",
            })

        async for p in db.payments.find(created):
            d = local_day(p.get("created_at"))
            if not in_range(d):
                continue
            is_c = p.get("entity_type") == "customer"
            name = (customers if is_c else suppliers).get(p.get("entity_id"), "")
            amt = r2(p.get("amount"))
            out.append({
                "key": f"payment:{p['_id']}", "source": "payment", "source_id": str(p["_id"]), "date": d,
                "time": p.get("created_at", ""), "type": "Arkëtim borxhi" if is_c else "Pagesë furnitori",
                "description": f"{name}{' · ' + p['notes'] if p.get('notes') else ''}", "invoice_no": "",
                "hyrje": amt if is_c else 0.0, "dalje": 0.0 if is_c else amt, "method": "bank" if p.get("bank") else "cash", "bank": p.get("bank"),
                "revenue": 0.0, "cost": 0.0, "cost_missing": False, "opex": 0.0, "value": amt, "note": "",
            })

        async for p in db.purchases.find(created):
            d = local_day(p.get("created_at"))
            if not in_range(d):
                continue
            seller = (p.get("citizen") or {}).get("name") if p.get("seller_type") == "citizen" else suppliers.get(p.get("supplier_id") or "", "")
            names = ", ".join(i.get("product_name", "") for i in p.get("items", []) or [])
            due = r2(p.get("due_amount"))
            out.append({
                "key": f"purchase:{p['_id']}", "source": "purchase", "source_id": str(p["_id"]), "date": d,
                "time": p.get("created_at", ""), "type": "Blerje", "description": f"{names} · {seller or 'Pa furnitor'}",
                "invoice_no": f"BL-{str(p['_id'])[-6:].upper()}", "hyrje": 0.0, "dalje": r2(p.get("paid_amount")),
                "method": "bank" if p.get("bank") else "cash", "bank": p.get("bank"), "revenue": 0.0, "cost": 0.0, "cost_missing": False, "opex": 0.0,
                "value": r2(p.get("total")), "note": f"Borxh {due:.2f} €" if due > 0.004 else "",
            })

        async for e in db.expenses.find({"date": {"$gte": d_from, "$lte": d_to}}):
            amt = r2(e.get("amount"))
            desc = e.get("category", "")
            if e.get("description"):
                desc += f" · {e['description']}"
            out.append({
                "key": f"expense:{e['_id']}", "source": "expense", "source_id": str(e["_id"]), "date": e["date"],
                "time": e.get("created_at", ""), "type": "Shpenzim", "description": desc, "invoice_no": "",
                "hyrje": 0.0, "dalje": amt, "method": "bank" if e.get("bank") else "cash", "bank": e.get("bank"),
                "revenue": 0.0, "cost": 0.0, "cost_missing": False, "opex": amt, "value": amt, "note": "",
            })

        async for rt in db.returns.find(created):
            d = local_day(rt.get("created_at"))
            if not in_range(d):
                continue
            sale = sales_by_id.get(rt.get("sale_id")) or (await db.sales.find_one({"_id": ObjectId(rt["sale_id"])}) if ObjectId.is_valid(rt.get("sale_id", "")) else None) or {}
            costs = {(i.get("product_id"), i.get("imei")): float(i.get("cost_price") or 0) for i in sale.get("items", []) or []}
            back_cost = sum(costs.get((i.get("product_id"), i.get("imei")), 0) * int(i.get("quantity") or 0) for i in rt.get("items", []) or [])
            names = ", ".join(i.get("product_name", "") for i in rt.get("items", []) or [])
            out.append({
                "key": f"return:{rt['_id']}", "source": "return", "source_id": str(rt["_id"]), "date": d,
                "time": rt.get("created_at", ""), "type": "Kthim", "description": f"{names} · {rt.get('customer_name') or 'Qytetar'}",
                "invoice_no": rt.get("invoice_no", ""), "hyrje": 0.0, "dalje": r2(rt.get("cash_refund")), "method": "cash",
                "revenue": -r2(rt.get("refund_total")), "cost": -r2(back_cost), "cost_missing": False, "opex": 0.0,
                "value": r2(rt.get("refund_total")), "note": "",
            })

        async for sv in db.services.find({"delivered_date": {"$gte": d_from, "$lte": d_to}, "status": "delivered", "deleted": {"$ne": True}}):
            paid = r2(sv.get("paid_amount"))
            cost = r2(r2(sv.get("parts_cost")) + sum(x["cost_price"] * x["quantity"] for x in sv.get("parts", [])))
            if paid <= 0 and cost <= 0:
                continue
            out.append({
                "key": f"service:{sv['_id']}", "source": "service", "source_id": str(sv["_id"]), "date": sv["delivered_date"],
                "time": sv.get("delivered_at") or sv.get("delivered_date", ""), "type": "Servis",
                "description": f"{sv.get('device', '')} · {sv.get('defect', '')} · {sv.get('first_name', '')} {sv.get('last_name', '')}".strip(),
                "invoice_no": sv.get("number", ""), "hyrje": paid, "dalje": 0.0, "method": sv.get("paid_method") or "cash",
                "revenue": paid, "cost": cost, "cost_missing": False, "opex": 0.0, "value": paid,
                "note": "falas / garanci" if paid <= 0 else ("pa kosto pjesësh" if sv.get("parts_cost") is None and not sv.get("parts") else ""),
            })

        async for t in db.bank_transfers.find({"date": {"$gte": d_from, "$lte": d_to}}):
            src, dst, amt = t["source"], t["target"], r2(t["amount"])
            lbl = lambda x: "Arka" if x == "cash" else x  # noqa: E731
            out.append({
                "key": f"transfer:{t['_id']}", "source": "transfer", "source_id": str(t["_id"]), "date": t["date"], "time": t["created_at"],
                "type": "Depozitim në bankë" if src == "cash" else ("Tërheqje nga banka" if dst == "cash" else "Transfertë bankare"),
                "description": f"{lbl(src)} → {lbl(dst)}{' · ' + t['note'] if t.get('note') else ''}", "invoice_no": "",
                "hyrje": amt if dst == "cash" else 0.0, "dalje": amt if src == "cash" else 0.0, "method": "cash" if "cash" in (src, dst) else "bank",
                "cash_in": amt if dst == "cash" else 0.0, "bank_in": amt if dst != "cash" else 0.0,
                "revenue": 0.0, "cost": 0.0, "cost_missing": False, "opex": 0.0, "value": amt, "note": "",
            })

        async for m in db.daybook_manual.find({"date": {"$gte": d_from, "$lte": d_to}, "deleted": {"$ne": True}}):
            amt = r2(m.get("amount"))
            is_in = m["kind"] == "in"
            out.append({
                "key": f"manual:{m['_id']}", "source": "manual", "source_id": str(m["_id"]), "date": m["date"],
                "time": m.get("created_at", ""), "type": "Hyrje manuale" if is_in else "Dalje manuale",
                "description": m.get("description", ""), "invoice_no": m.get("invoice_no", ""),
                "hyrje": amt if is_in else 0.0, "dalje": 0.0 if is_in else amt, "method": m.get("method", "cash"),
                "revenue": amt if (is_in and m.get("in_profit")) else 0.0, "cost": 0.0, "cost_missing": False,
                "opex": amt if (not is_in and m.get("in_profit")) else 0.0, "value": amt, "in_profit": bool(m.get("in_profit")),
                "kind": m["kind"], "note": "",
            })

        for e in out:
            o = overrides.get(e["key"])
            e["corrected"] = bool(o)
            if o:
                for f in OVERRIDABLE:
                    if o.get(f) is not None:
                        if f == "cost":
                            e["cost"], e["cost_missing"] = r2(o["cost"]), False
                        else:
                            e[f] = r2(o[f]) if f in ("hyrje", "dalje") else o[f]
            if o and (o.get("hyrje") is not None or o.get("method") is not None) or "cash_in" not in e:
                # flows follow the (possibly corrected) amount and method
                to_bank = e["method"] == "bank"
                e["cash_in"], e["bank_in"] = (0.0, e["hyrje"]) if to_bank else (e["hyrje"], 0.0)
                if e["method"] == "mixed":
                    e["method"] = "cash"
            if e["source"] == "transfer":
                e["cash_out"] = e["dalje"]
                e["bank_out"] = e["value"] if not e["dalje"] else 0.0
            else:
                e["cash_out"], e["bank_out"] = (0.0, e["dalje"]) if e["method"] == "bank" else (e["dalje"], 0.0)
            e.setdefault("vat", 0.0)
            e.setdefault("bank_fee", 0.0)
            e["profit"] = r2(e["revenue"] - e["vat"] - e["cost"] - e["bank_fee"] - e["opex"])
        out.sort(key=lambda e: (e["date"], e["time"]))
        return out

    def cash_net(e) -> float:
        return e["cash_in"] - e["cash_out"]

    def summarize(entries: list, tax_pct: float) -> dict:
        by = lambda src, field: r2(sum(e["cash_in" if field == "hyrje" else "cash_out"] for e in entries if e["source"] == src))  # noqa: E731
        sales = [e for e in entries if e["source"] == "sale"]
        t_cash = r2(sum(e["cash_in"] for e in sales))
        t_card = r2(sum(e["bank_in"] for e in sales))
        vat = r2(sum(e["vat"] for e in entries))
        fees = r2(sum(e["bank_fee"] for e in entries))
        revenue = r2(sum(e["revenue"] for e in entries if e["source"] in ("sale", "return")))
        cost = r2(sum(e["cost"] for e in entries if e["source"] in ("sale", "return")))
        service_rev = r2(sum(e["revenue"] for e in entries if e["source"] == "service"))
        service_cost = r2(sum(e["cost"] for e in entries if e["source"] == "service"))
        other_income = r2(sum(e["revenue"] for e in entries if e["source"] == "manual"))
        opex = r2(sum(e["opex"] for e in entries))
        gross = r2(revenue - vat - cost)
        pbt = r2(gross - fees + service_rev - service_cost + other_income - opex)
        tax = r2(max(0.0, pbt) * tax_pct / 100)
        return {
            "sales_count": len(sales), "turnover_cash": t_cash, "turnover_card": t_card, "turnover_total": r2(sum(e["revenue"] for e in sales)), "turnover_debt": r2(sum(e["revenue"] - e["cash_in"] - e["bank_in"] for e in sales)),
            "vat": vat, "bank_fees": fees, "profit_before_tax": pbt, "profit_tax": tax, "profit_tax_pct": tax_pct,
            "cash_sales": by("sale", "hyrje"),
            "debt_collections": r2(sum(e["cash_in"] for e in entries if e["type"] == "Arkëtim borxhi")),
            "manual_in": by("manual", "hyrje"),
            "service_cash": by("service", "hyrje"),
            "purchases_out": by("purchase", "dalje"),
            "expenses_out": by("expense", "dalje"),
            "supplier_payments": r2(sum(e["cash_out"] for e in entries if e["type"] == "Pagesë furnitori")),
            "returns_out": by("return", "dalje"),
            "manual_out": by("manual", "dalje"),
            "cash_in": r2(sum(e["cash_in"] for e in entries)),
            "cash_out": r2(sum(e["cash_out"] for e in entries)),
            "bank_in": r2(sum(e["bank_in"] for e in entries)),
            "bank_out": r2(sum(e["bank_out"] for e in entries)),
            "sales_total": r2(sum(e["revenue"] for e in entries if e["source"] == "sale")),
            "purchases_total": r2(sum(e["value"] for e in entries if e["source"] == "purchase")),
            "expenses_total": r2(sum(e["value"] for e in entries if e["source"] == "expense")),
            "returns_total": r2(sum(e["value"] for e in entries if e["source"] == "return")),
            "revenue": revenue, "cogs": cost, "gross_profit": gross, "other_income": other_income, "opex": opex,
            "service_income": service_rev, "service_parts_cost": service_cost, "service_profit": r2(service_rev - service_cost),
            "service_delivered": sum(1 for e in entries if e["source"] == "service"),
            "net_profit": r2(pbt - tax),
            "profit_incomplete": any(e["cost_missing"] for e in entries),
            "missing_cost_count": sum(1 for e in entries if e["cost_missing"]),
        }

    def count_total(counts: dict) -> float:
        return r2(sum(float(d) * max(0, int(counts.get(d, 0) or 0)) for d in DENOMS))

    async def ledger(d_from: str, d_to: str):
        """Entries from the cash start date to d_to, plus the opening cash at d_from."""
        st = await settings()
        start = st["start_date"] or d_from
        lo = min(start, d_from)
        entries = await collect(lo, d_to)
        before = [e for e in entries if start <= e["date"] < d_from]
        opening = r2(st["opening_cash"] + sum(cash_net(e) for e in before)) if d_from >= start else None
        return st, start, entries, opening

    @router.get("/day")
    async def day(date: str, _a: dict = Depends(require_admin)):
        parse_day(date)
        st, start, entries, opening = await ledger(date, date)
        todays = [e for e in entries if e["date"] == date]
        s = summarize(todays, (await get_rates(db))["profit_tax_pct"])
        closing = r2(opening + s["cash_in"] - s["cash_out"]) if opening is not None else None
        c = await db.cash_counts.find_one({"date": date})
        count = None
        if c:
            counted = count_total(c.get("counts", {}))
            count = {"counts": c.get("counts", {}), "total": counted, "note": c.get("note", ""),
                     "difference": r2(counted - closing) if closing is not None else None,
                     "updated_at": c.get("updated_at"), "updated_by": c.get("updated_by_name", "")}
        received = [{"id": str(x["_id"]), "number": x.get("number"), "name": f"{x.get('first_name', '')} {x.get('last_name', '')}".strip(),
                     "phone": x.get("phone"), "device": x.get("device"), "defect": x.get("defect"), "price": x.get("price"),
                     "status": x.get("status"), "paid_amount": x.get("paid_amount")}
                    async for x in db.services.find({"received_date": date, "deleted": {"$ne": True}}).sort("created_at", 1)]
        return {"date": date, "services_received": received, "settings": st, "before_start": opening is None, "entries": todays, "summary": s,
                "opening_cash": opening, "closing_cash": closing, "count": count, "turnover": s["sales_total"]}

    @router.get("/settings")
    async def get_settings(_a: dict = Depends(require_admin)):
        return await settings()

    @router.put("/settings")
    async def put_settings(body: SettingsIn, user: dict = Depends(require_admin)):
        parse_day(body.start_date)
        before = await settings()
        await db.daybook_settings.update_one({"_id": "cash"}, {"$set": {"start_date": body.start_date, "opening_cash": r2(body.opening_cash)}}, upsert=True)
        await audit("settings", "update", before, body.model_dump(), user)
        return await settings()

    @router.post("/manual")
    async def add_manual(body: ManualIn, user: dict = Depends(require_admin)):
        parse_day(body.date)
        if body.amount <= 0:
            raise HTTPException(400, "Shuma duhet të jetë më e madhe se 0")
        doc = {**body.model_dump(), "amount": r2(body.amount), "description": body.description.strip()[:300],
               "created_at": now_utc().isoformat(), "created_by": user["id"]}
        r = await db.daybook_manual.insert_one(doc)
        await audit(f"manual:{r.inserted_id}", "create", None, body.model_dump(), user)
        return {"id": str(r.inserted_id)}

    @router.put("/manual/{mid}")
    async def edit_manual(mid: str, body: ManualIn, user: dict = Depends(require_admin)):
        parse_day(body.date)
        if body.amount <= 0:
            raise HTTPException(400, "Shuma duhet të jetë më e madhe se 0")
        old = await db.daybook_manual.find_one({"_id": ObjectId(mid), "deleted": {"$ne": True}}) if ObjectId.is_valid(mid) else None
        if not old:
            raise HTTPException(404, "Nuk u gjet")
        new = {**body.model_dump(), "amount": r2(body.amount)}
        await db.daybook_manual.update_one({"_id": old["_id"]}, {"$set": new})
        await audit(f"manual:{mid}", "update", {k: old.get(k) for k in new}, new, user)
        return {"ok": True}

    @router.delete("/manual/{mid}")
    async def delete_manual(mid: str, body: DeleteIn, user: dict = Depends(require_admin)):
        if not body.reason.strip():
            raise HTTPException(400, "Shkruaj arsyen e fshirjes")
        old = await db.daybook_manual.find_one({"_id": ObjectId(mid), "deleted": {"$ne": True}}) if ObjectId.is_valid(mid) else None
        if not old:
            raise HTTPException(404, "Nuk u gjet")
        # soft delete: the record stays in the database and in the history
        await db.daybook_manual.update_one({"_id": old["_id"]}, {"$set": {"deleted": True, "deleted_at": now_utc().isoformat(), "deleted_by": user["id"]}})
        await audit(f"manual:{mid}", "delete", {k: old.get(k) for k in ("date", "kind", "amount", "method", "description")}, None, user, body.reason)
        return {"ok": True}

    @router.put("/override")
    async def put_override(body: OverrideIn, user: dict = Depends(require_admin)):
        src = body.key.split(":")[0]
        if src not in ("sale", "purchase", "expense", "payment", "return"):
            raise HTTPException(400, "Ky transaksion nuk mund të korrigjohet këtu")
        coll = {"sale": db.sales, "purchase": db.purchases, "expense": db.expenses, "payment": db.payments, "return": db.returns}[src]
        sid = body.key.partition(":")[2]
        if not ObjectId.is_valid(sid) or not await coll.find_one({"_id": ObjectId(sid)}, {"_id": 1}):
            raise HTTPException(404, "Transaksioni nuk u gjet")
        fields = {f: getattr(body, f) for f in OVERRIDABLE if getattr(body, f) is not None}
        for f in ("hyrje", "dalje", "cost"):
            if f in fields:
                if fields[f] < 0:
                    raise HTTPException(400, "Shuma nuk mund të jetë negative")
                fields[f] = r2(fields[f])
        old = await db.daybook_overrides.find_one({"key": body.key})
        await db.daybook_overrides.update_one({"key": body.key}, {"$set": {**fields, "key": body.key, "updated_at": now_utc().isoformat()}}, upsert=True)
        await audit(body.key, "correct", {f: (old or {}).get(f) for f in fields}, fields, user, body.reason)
        return {"ok": True}

    @router.get("/entry")
    async def get_entry(key: str, _a: dict = Depends(require_admin)):
        src, _, sid = key.partition(":")
        coll = {"sale": db.sales, "purchase": db.purchases, "expense": db.expenses, "payment": db.payments,
                "return": db.returns, "manual": db.daybook_manual}.get(src)
        doc = await coll.find_one({"_id": ObjectId(sid)}) if coll is not None and ObjectId.is_valid(sid) else None
        if not doc:
            raise HTTPException(404, "Transaksioni nuk u gjet")
        d = doc.get("document_date") or doc.get("date") or local_day(doc.get("created_at"))
        found = next((e for e in await collect(d, d) if e["key"] == key), None)
        if not found:
            raise HTTPException(404, "Transaksioni nuk u gjet")
        return found

    @router.delete("/override")
    async def delete_override(key: str, user: dict = Depends(require_admin)):
        old = await db.daybook_overrides.find_one_and_delete({"key": key})
        if old:
            await audit(key, "restore", {f: old.get(f) for f in OVERRIDABLE if old.get(f) is not None}, None, user)
        return {"ok": True}

    @router.get("/audit")
    async def get_audit(key: Optional[str] = None, _a: dict = Depends(require_admin)):
        q = {"key": key} if key else {}
        rows = await db.daybook_audit.find(q).sort("at", -1).to_list(500)
        return [{**{k: v for k, v in r.items() if k != "_id"}, "id": str(r["_id"])} for r in rows]

    @router.put("/count/{date}")
    async def put_count(date: str, body: CountIn, user: dict = Depends(require_admin)):
        parse_day(date)
        counts = {d: max(0, int(body.counts.get(d, 0) or 0)) for d in DENOMS}
        old = await db.cash_counts.find_one({"date": date})
        await db.cash_counts.update_one({"date": date}, {"$set": {
            "counts": counts, "total": count_total(counts), "note": body.note[:300],
            "updated_at": now_utc().isoformat(), "updated_by": user["id"], "updated_by_name": user.get("name", ""),
        }}, upsert=True)
        await audit(f"count:{date}", "update" if old else "create", (old or {}).get("counts"), counts, user)
        return {"ok": True, "total": count_total(counts)}

    async def report_rows(d_from: str, d_to: str, group: str):
        if parse_day(d_from) > parse_day(d_to):
            raise HTTPException(400, "Data 'nga' duhet të jetë para datës 'deri'")
        st, start, entries, _ = await ledger(d_from, d_to)
        tax_pct = (await get_rates(db))["profit_tax_pct"]
        counts = {c["date"]: count_total(c.get("counts", {})) async for c in db.cash_counts.find({"date": {"$gte": d_from, "$lte": d_to}})}
        by_day: dict = {}
        for e in entries:
            by_day.setdefault(e["date"], []).append(e)
        cash = st["opening_cash"]
        cur = parse_day(min(start, d_from))
        end = parse_day(d_to)
        groups: dict = {}
        while cur <= end:
            d = cur.isoformat()
            day_entries = by_day.get(d, [])
            if d >= start:
                cash = r2(cash + sum(cash_net(e) for e in day_entries))
            if d >= d_from:
                k = d if group == "day" else d[:7] if group == "month" else d[:4]
                g = groups.setdefault(k, {"period": k, "entries": [], "cash_end": None, "count_difference": 0.0, "counted_days": 0})
                g["entries"] += day_entries
                g["cash_end"] = cash if d >= start else None
                if d in counts and d >= start:
                    g["count_difference"] = r2(g["count_difference"] + counts[d] - cash)
                    g["counted_days"] += 1
            cur += timedelta(days=1)
        rows = []
        for g in groups.values():
            s = summarize(g.pop("entries"), tax_pct)
            rows.append({**g, **s})
        keys = ("sales_count", "turnover_cash", "turnover_card", "turnover_total", "sales_total", "service_income", "purchases_total",
                "expenses_total", "returns_total", "cash_in", "cash_out", "bank_in", "bank_out", "revenue", "vat", "bank_fees", "cogs",
                "gross_profit", "opex", "profit_before_tax", "count_difference")
        totals = {k: r2(sum(r[k] for r in rows)) for k in keys}
        # profit tax on the whole period (losses of one day offset profits of another)
        totals["profit_tax"] = r2(max(0.0, totals["profit_before_tax"]) * tax_pct / 100)
        totals["net_profit"] = r2(totals["profit_before_tax"] - totals["profit_tax"])
        totals["profit_incomplete"] = any(r["profit_incomplete"] for r in rows)
        totals["cash_end"] = rows[-1]["cash_end"] if rows else None
        return rows, totals

    @router.get("/report")
    async def report(date_from: str, date_to: str, group: Literal["day", "month", "year"] = "day", _a: dict = Depends(require_admin)):
        rows, totals = await report_rows(date_from, date_to, group)
        return {"rows": rows, "totals": totals}

    COLS = [("period", "Periudha"), ("sales_count", "Nr. shitjeve"), ("turnover_cash", "Qarkullimi Cash"), ("turnover_card", "Qarkullimi Card-Bank"),
            ("turnover_total", "Qarkullimi total"), ("purchases_total", "Blerjet"), ("expenses_total", "Shpenzimet"), ("returns_total", "Kthimet"),
            ("service_income", "Servisi"), ("bank_fees", "Komisione bankare"), ("vat", "TVSH"), ("cogs", "Kosto e mallit"),
            ("profit_before_tax", "Fitimi para tatimit"), ("profit_tax", "Tatimi në fitim"), ("net_profit", "Fitimi i pastër"),
            ("cash_in", "Hyrje Cash"), ("cash_out", "Dalje Cash"), ("cash_end", "Arka në fund"), ("count_difference", "Dif. numërimit")]

    @router.get("/report.xlsx")
    async def report_xlsx(date_from: str, date_to: str, group: Literal["day", "month", "year"] = "day", _a: dict = Depends(require_admin)):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        rows, totals = await report_rows(date_from, date_to, group)
        wb = Workbook()
        ws = wb.active
        ws.title = "Qarkullimet"
        ws.append([f"Qarkullimet Ditore · {date_from} – {date_to}"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([c[1] for c in COLS])
        for cell in ws[2]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
        for r in rows:
            ws.append([r[k] if r[k] is not None else "" for k, _ in COLS])
        ws.append(["TOTALI"] + [totals.get(k) if totals.get(k) is not None else "" for k, _ in COLS[1:]])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        for row in ws.iter_rows(min_row=3, min_col=2):
            for cell in row:
                if cell.column != 2:
                    cell.number_format = '#,##0.00 "€"'
        from openpyxl.utils import get_column_letter
        for i in range(1, len(COLS) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 16
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f'attachment; filename="qarkullimet_{date_from}_{date_to}.xlsx"'})

    @router.get("/report.pdf")
    async def report_pdf(date_from: str, date_to: str, group: Literal["day", "month", "year"] = "day", _a: dict = Depends(require_admin)):
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import mm
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        rows, totals = await report_rows(date_from, date_to, group)
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm)
        fmt = lambda v: "—" if v is None else (f"{v:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".") if isinstance(v, (int, float)) else str(v))  # noqa: E731
        from reportlab.lib.styles import ParagraphStyle
        hs = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=6, leading=7, textColor=colors.white, alignment=1)
        val = lambda k, v: str(int(v)) if k == "sales_count" and v is not None else fmt(v)  # noqa: E731
        data = ([[Paragraph(c[1], hs) for c in COLS]] + [[val(k, r[k]) if k != "period" else r[k] for k, _ in COLS] for r in rows]
                + [["TOTALI"] + [val(k, totals.get(k)) for k, _ in COLS[1:]]])
        t = Table(data, repeatRows=1, colWidths=[18 * mm] + [(277 - 18) / (len(COLS) - 1) * mm] * (len(COLS) - 1))
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F766E")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 6), ("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F1F5F9")]),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E2E8F0")),
        ]))
        title = getSampleStyleSheet()["Heading2"]
        story = company_header(f"Raporti i qarkullimeve · {date_from} – {date_to}") + [Paragraph("Qarkullimet Ditore", title), Spacer(1, 4), t]
        if totals["profit_incomplete"]:
            story += [Spacer(1, 6), Paragraph("Kujdes: disa produkte nuk kanë çmim blerjeje – fitimi është i paplotë.", getSampleStyleSheet()["Normal"])]
        doc.build(story)
        buf.seek(0)
        return StreamingResponse(buf, media_type="application/pdf",
                                 headers={"Content-Disposition": f'inline; filename="qarkullimet_{date_from}_{date_to}.pdf"'})

    return router
